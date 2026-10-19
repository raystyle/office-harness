"""Probe: does an officecli round-trip refresh openpyxl formula caches."""

from __future__ import annotations

import subprocess
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook

d = Path(sys.argv[1])
f = d / "recalc-probe.xlsx"
wb = Workbook()
ws = wb.active
ws["A1"] = 1
ws["A2"] = 2
ws["A3"] = "=SUM(A1:A2)"  # openpyxl writes NO cached value for this
wb.save(str(f))


def cached() -> object:
    """Read the cached value of A3 through the data_only lens."""
    wbr = load_workbook(str(f), data_only=True)
    return wbr.active["A3"].value


print(f"before officecli round-trip: A3 cached = {cached()!r}")
oc = r"D:\tmp\oc162-review\officecli.exe"
r = subprocess.run(
    [oc, "set", str(f), "/Sheet/A3", "--prop", "formula==SUM(A1:A2)", "--json"],
    capture_output=True, text=True, timeout=60,
)
print(f"  set formula: rc={r.returncode} {(r.stdout or r.stderr)[:120]!r}")
r = subprocess.run([oc, "close", str(f), "--json"], capture_output=True, text=True, timeout=60)
print(f"  close: rc={r.returncode}")
print(f"after officecli set:          A3 cached = {cached()!r}")
