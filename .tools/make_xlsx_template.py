"""Make the xlsx seed template: anchors + a list slot, openpyxl-authored."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

HERE = Path(__file__).resolve().parents[1] / "templates" / "xlsx"


def main() -> None:
    """Write the xlsx seed template (run from repo root via uv)."""
    HERE.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "report"
    ws["A1"] = "{{ title }}"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "作者:{{ author }}   日期:{{ date }}"
    ws["A4"] = "项目"
    ws["B4"] = "结果"
    for cell in ("A4", "B4"):
        ws[cell].font = Font(bold=True)
        ws[cell].alignment = Alignment(horizontal="center")
    ws["A5"] = "{{ rows }}"
    ws["B5"] = "{{ notes }}"
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 40
    out = HERE / "report.xlsx"
    wb.save(str(out))
    print(f"written: {out}")


if __name__ == "__main__":
    main()
