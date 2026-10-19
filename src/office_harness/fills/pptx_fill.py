"""pptx fill via python-pptx: {{key}} anchors rewritten over whole paragraphs.

Design stance (industry lesson): jinja-in-pptx dies on PowerPoint's run
splitting, so our templates own the runs — one ``{{key}}`` per shape. The
rewriter joins each paragraph's runs, substitutes, then writes the result
into the first run and blanks the rest, immune to how Word/PPT split runs.

TODO(M4): ``{{for:name}}`` slide duplication per data row — python-pptx has
no public slide-copy API; lands with its own tests.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from pptx import Presentation

_ANCHOR = re.compile(r"\{\{\s*([a-zA-Z_][a-zA-Z0-9_.]*)\s*\}\}")


def _iter_text_frames(shape: Any) -> list[Any]:
    """All text frames of one shape, descending into groups."""
    frames: list[Any] = []
    if shape.has_text_frame:
        frames.append(shape.text_frame)
    if getattr(shape, "shapes", None):
        for sub in shape.shapes:
            frames.extend(_iter_text_frames(sub))
    return frames


def _apply_anchors(text_frame: Any, context: dict[str, Any]) -> int:
    """Rewrite {{key}} anchors in one text frame; returns the hit count."""
    hits = 0

    def lookup(match: re.Match[str]) -> str:
        nonlocal hits
        key = match.group(1)
        value: Any = context
        for part in key.split("."):
            if isinstance(value, dict) and part in value:
                value = value[part]
            else:
                return match.group(0)  # unresolved anchors stay literal
        hits += 1
        return str(value)

    for para in text_frame.paragraphs:
        if not para.runs:
            continue
        full = "".join(run.text for run in para.runs)
        if not _ANCHOR.search(full):
            continue
        replaced = _ANCHOR.sub(lookup, full)
        para.runs[0].text = replaced
        for run in para.runs[1:]:
            run.text = ""
    return hits


def fill(template: str, data: dict[str, Any], out: str) -> dict[str, Any]:
    """Fill one pptx template; returns a report dict."""
    prs = Presentation(template)
    anchors = 0
    for slide in prs.slides:
        for shape in slide.shapes:
            for tf in _iter_text_frames(shape):
                anchors += _apply_anchors(tf, data)
    out_path = Path(out)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    return {
        "engine": "python-pptx",
        "template": template,
        "out": str(out_path),
        "anchors": anchors,
        "bytes": out_path.stat().st_size,
    }
