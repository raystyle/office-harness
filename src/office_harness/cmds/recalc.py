"""recalc: refresh xlsx formula caches by rewriting formulas through officecli.

openpyxl writes formulas with NO cached values (data_only reads return None).
officecli evaluates formulas on every write, so rewriting each formula cell
via the engine refreshes the cache with no LibreOffice detour (probe evidence
in .tools/probe_recalc.py, 2026-10-19).
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from office_harness import envelope
from office_harness.cmds.common import OpError, parse_opts
from office_harness.engine import officecli
from office_harness.fills.xlsx_fill import _formula_lint

_OPTS: dict[str, str | None] = {"--yes": None}


def _collect_formulas(file: str) -> tuple[list[dict[str, str]], list[str]]:
    """Return formula cells and compatibility warnings from one workbook."""
    wb = load_workbook(file)
    cells: list[dict[str, str]] = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    cells.append({"sheet": ws.title, "ref": cell.coordinate, "formula": cell.value})
    return cells, _formula_lint(wb)


def run(g: envelope.Globals, tokens: list[str]) -> int:
    """Recalculate formula caches in place; dry-run lists cells until --yes."""
    opts, positional = parse_opts(tokens, _OPTS)
    if len(positional) != 1:
        raise OpError("usage", "recalc takes exactly one .xlsx path")
    file = positional[0]
    if not Path(file).exists():
        raise OpError("file_not_found", f"file not found: {file}")
    if Path(file).suffix.lower() not in (".xlsx", ".xlsm"):
        raise OpError("unsupported_format", f"recalc supports xlsx/xlsm, got {Path(file).suffix}")

    cells, warnings = _collect_formulas(file)

    if "--yes" not in opts:
        report: dict[str, Any] = {
            "dryrun": True,
            "file": file,
            "formula_cells": len(cells),
            "cells": cells[:50],
            "formula_warnings": warnings,
        }
        def human(d: dict[str, Any]) -> None:
            print(f"dry-run: {d['formula_cells']} formula cell(s); --yes to recalc")

        return envelope.finish(g, report, human=human)

    refreshed = 0
    failed: list[str] = []
    for c in cells:
        path = f"/{c['sheet']}/{c['ref']}"
        payload = officecli.set_formula(file, path, c["formula"])
        if payload.get("ok"):
            refreshed += 1
        else:
            failed.append(f"{path}: {payload.get('error', {}).get('error', 'unknown')}")
    officecli.close_quietly(file)

    wbr = load_workbook(file, data_only=True)
    cached: dict[str, Any] = {}
    for c in cells:
        v = wbr[c["sheet"]][c["ref"]].value
        if v is not None:
            cached[f"{c['sheet']}!{c['ref']}"] = v

    report = {
        "dryrun": False,
        "file": file,
        "formula_cells": len(cells),
        "refreshed": refreshed,
        "cached_now": len(cached),
        "cache_values": cached,
        "formula_warnings": warnings,
    }
    if failed:
        report["failed"] = failed
    return envelope.finish(
        g,
        report,
        human=lambda d: print(
            f"recalced {d['refreshed']}/{d['formula_cells']} formulas, cache live on {d['cached_now']}"
        ),
    )

