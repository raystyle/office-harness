"""Library fill layer tests over real template files (no engine involved)."""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook
from pptx import Presentation

from office_harness import run

TEMPLATES = Path(__file__).resolve().parents[2] / "templates"


def test_fill_docx_real_template(tmp_path: Path, capsys) -> None:
    data = {
        "title": "office-harness 周报",
        "author": "claude-2",
        "date": "2026-10-19",
        "summary": "M3 docx 线验收。",
        "tags": ["M1", "M2", "M3"],
    }
    dfile = tmp_path / "d.json"
    dfile.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    out = tmp_path / "o.docx"
    code = run.main(
        ["fill", str(TEMPLATES / "docx" / "report.docx"), "--data", str(dfile),
         "--out", str(out), "--yes", "--json"]
    )
    payload = json.loads(capsys.readouterr().out)
    assert code == 0
    assert payload["data"]["engine"] == "docxtpl"
    assert out.exists() and out.stat().st_size > 0


def test_fill_pptx_real_template(tmp_path: Path, capsys) -> None:
    data = {"title": "office-mesh 结网汇报", "speaker": "lan-win-claude-2", "date": "2026-10-19"}
    dfile = tmp_path / "d.json"
    dfile.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    out = tmp_path / "o.pptx"
    code = run.main(
        ["fill", str(TEMPLATES / "pptx" / "deck.pptx"), "--data", str(dfile),
         "--out", str(out), "--yes", "--json"]
    )
    payload = json.loads(capsys.readouterr().out)
    assert code == 0
    assert payload["data"]["anchors"] >= 2
    prs = Presentation(str(out))
    texts = [s.text_frame.text for s in prs.slides[0].shapes if s.has_text_frame]
    joined = "\n".join(texts)
    assert "office-mesh 结网汇报" in joined
    assert "lan-win-claude-2" in joined


def test_fill_xlsx_anchor_and_rows(tmp_path: Path, capsys) -> None:
    tpl = tmp_path / "t.xlsx"
    wb = __import__("openpyxl").Workbook()
    ws = wb.active
    ws["A1"] = "{{ name }}"
    ws["B1"] = "{{ visits }}"
    ws["A2"] = "{{ series }}"
    wb.save(str(tpl))
    data = {"name": "office-harness", "visits": 42, "series": [10, 20, 30]}
    dfile = tmp_path / "d.json"
    dfile.write_text(json.dumps(data), encoding="utf-8")
    out = tmp_path / "o.xlsx"
    code = run.main(["fill", str(tpl), "--data", str(dfile), "--out", str(out), "--yes", "--json"])
    capsys.readouterr()
    assert code == 0
    wbo = load_workbook(str(out))
    wso = wbo.active
    assert wso["A1"].value == "office-harness"
    assert wso["B1"].value == 42
    assert [wso.cell(row=2 + i, column=1).value for i in range(3)] == [10, 20, 30]


def test_fill_pptx_for_slide_duplication(tmp_path: Path, capsys) -> None:
    from pptx import Presentation as Build
    from pptx.util import Inches

    tpl = tmp_path / "loop.pptx"
    prs = Build()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box1 = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    box1.text_frame.text = "{{ title }}"
    box2 = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1))
    box2.text_frame.text = "{{ name }} - {{ note }}"
    slide.notes_slide.notes_text_frame.text = "{{for:modules}}"
    prs.save(str(tpl))

    data = {
        "title": "模块巡礼",
        "modules": [
            {"name": "render", "note": "texenv"},
            {"name": "compile", "note": "tectonic"},
            {"name": "fill", "note": "四族"},
        ],
    }
    dfile = tmp_path / "d.json"
    dfile.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
    out = tmp_path / "o.pptx"
    code = run.main(["fill", str(tpl), "--data", str(dfile), "--out", str(out), "--yes", "--json"])
    payload = json.loads(capsys.readouterr().out)
    assert code == 0
    assert payload["data"]["clones"] == 3
    prs_out = Presentation(str(out))
    assert len(prs_out.slides._sldIdLst) == 4  # noqa: SLF001 - source + 3 clones
    texts: list[str] = []
    for s in prs_out.slides:
        for shape in s.shapes:
            if shape.has_text_frame:
                texts.append(shape.text_frame.text)
    assert "模块巡礼" in texts
    assert "render - texenv" in texts
    assert "compile - tectonic" in texts
    assert "fill - 四族" in texts
    # for-marker stripped from the source notes
    assert "{{for:" not in (prs_out.slides[0].notes_slide.notes_text_frame.text or "")


def test_fill_reports_unresolved_and_formula_lint(tmp_path: Path, capsys) -> None:
    tpl = tmp_path / "t.xlsx"
    wb = __import__("openpyxl").Workbook()
    ws = wb.active
    ws["A1"] = "{{ known }}"
    ws["A2"] = "{{ missing_key }}"
    ws["B1"] = "=XLOOKUP(1, A:A, B:B)"
    wb.save(str(tpl))
    data = {"known": "v"}
    dfile = tmp_path / "d.json"
    dfile.write_text(json.dumps(data), encoding="utf-8")
    out = tmp_path / "o.xlsx"
    code = run.main(["fill", str(tpl), "--data", str(dfile), "--out", str(out), "--yes", "--json"])
    payload = json.loads(capsys.readouterr().out)
    assert code == 0
    assert payload["data"]["unresolved"] == ["missing_key"]
    assert any("XLOOKUP" in w for w in payload["data"]["formula_warnings"])


def test_fill_pptx_reports_unresolved(tmp_path: Path, capsys) -> None:
    from pptx import Presentation as Build
    from pptx.util import Inches

    tpl = tmp_path / "u.pptx"
    prs = Build()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    box.text_frame.text = "{{ nope }}"
    prs.save(str(tpl))
    dfile = tmp_path / "d.json"
    dfile.write_text(json.dumps({"x": 1}), encoding="utf-8")
    out = tmp_path / "o.pptx"
    code = run.main(["fill", str(tpl), "--data", str(dfile), "--out", str(out), "--yes", "--json"])
    payload = json.loads(capsys.readouterr().out)
    assert code == 0
    assert payload["data"]["unresolved"] == ["nope"]


def test_fill_dryrun_then_yes(tmp_path: Path, capsys) -> None:
    data = {"title": "t", "author": "a", "date": "d", "summary": "s", "tags": ["x"]}
    dfile = tmp_path / "d.json"
    dfile.write_text(json.dumps(data), encoding="utf-8")
    tpl = str(TEMPLATES / "docx" / "report.docx")
    code = run.main(["fill", tpl, "--data", str(dfile), "--json", "--full-output"])
    payload = json.loads(capsys.readouterr().out)
    assert code == 0
    assert payload["data"]["dryrun"] is True
    assert payload["meta"]["cta"]["commands"][0]["command"].endswith("--yes")

