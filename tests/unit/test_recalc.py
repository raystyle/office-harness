"""recalc command tests: dry-run face, fake-engine apply, cache assertions."""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook

from office_harness import run


def _make_wb(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws["A1"] = 1
    ws["A2"] = 2
    ws["A3"] = "=SUM(A1:A2)"
    wb.save(str(path))


def test_recalc_dryrun_lists_cells(tmp_path: Path, capsys) -> None:
    f = tmp_path / "t.xlsx"
    _make_wb(f)
    code = run.main(["recalc", str(f), "--json"])
    payload = json.loads(capsys.readouterr().out)
    assert code == 0
    assert payload["data"]["dryrun"] is True
    assert payload["data"]["formula_cells"] == 1
    assert payload["data"]["cells"][0]["formula"] == "=SUM(A1:A2)"


def test_recalc_apply_via_fake_engine(fake_engine, tmp_path: Path, capsys) -> None:
    f = tmp_path / "t.xlsx"
    _make_wb(f)
    fake_engine.script(("officecli", "set"), (0, '{"ok": true, "data": "updated"}', ""))
    fake_engine.script(("officecli", "close"), (0, '{"ok": true, "data": "closed"}', ""))
    code = run.main(["recalc", str(f), "--yes", "--json", "--full-output"])
    payload = json.loads(capsys.readouterr().out)
    assert code == 0
    assert payload["data"]["refreshed"] == 1
    # the fake engine cannot write caches, so the live-cache count stays 0
    # here; the live path is covered by the real-engine smoke (.tools/probe)
    assert payload["data"]["cached_now"] == 0
    set_calls = [c for c in fake_engine.calls if c[1] == "set"]
    assert set_calls[0][3] == "/Sheet/A3"
    assert set_calls[0][5].startswith("formula==SUM")
    assert payload["meta"]["command"] == "office-harness recalc"


def test_recalc_rejects_non_xlsx(tmp_path: Path, capsys) -> None:
    f = tmp_path / "t.docx"
    f.write_bytes(b"zz")
    code = run.main(["recalc", str(f)])
    payload = json.loads(capsys.readouterr().err.strip().splitlines()[-1])
    assert code == 2
    assert payload["code"] == "unsupported_format"
